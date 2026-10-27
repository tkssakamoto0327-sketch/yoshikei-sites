#!/usr/bin/env python3
"""ヨシケイキッチン レシピ帳：週間レシピの Excel → recipe/data.js に取り込む

usage: python3 -I tools/recipe_xlsx2data.py IN.xlsx recipe/data.js [--check]

- Excel はシート1枚＝1日（シート名「11月23日」、B2「メニュー日：2026年11月23日」）
- 各ブロックは見出し「（手作りおかず朝）」のように コース名＋朝/昼/夕
  - B列：料理名（その行の F列が作り方）
  - C列：材料、D列：1人分、E列：3人分、G列：アレルギー
  - 赤字（FFFF0000）の材料＝ご家庭で用意するもの → own:true
  - 「栄養価　…」の行がその食事の栄養価
- 写真は Excel に入っていない。献立表のPDF（kondate/pdf/<その週の月曜>_<コース>.pdf）に
  表の形（横＝日付、縦＝朝・昼・夕）で入っているので、位置で対応を決めて取り出し、recipe/photos/ に置く
  （横420px の JPEG。PyMuPDF が必要：pip install --break-system-packages pymupdf）
  献立表PDFが無いコースは、既存の写真があれば引き継ぎ、無ければ null（盛付写真なしで表示される）
- --check：書き込まずに、既存の data.js と内容が一致するかだけ確かめる（変換ルールの検証用）
"""
import sys, re, json, datetime
from fractions import Fraction
import openpyxl
import os, hashlib, tempfile
from PIL import Image

COURSE = [  # 長い名前から判定
    ('ﾎﾞﾘｭｰﾑ手作りおかず', 'volume'), ('ボリューム手作りおかず', 'volume'),
    ('超簡単プラス', 'chokantanplus'), ('超簡単おかず', 'chokantan'),
    ('簡単おかず', 'kantan'), ('手作りおかず', 'tezukuri'), ('手づくりおかず', 'tezukuri'),
]
MEAL = {'朝': '朝食', '昼': '昼食', '夕': '夕食'}
WD = '月火水木金土日'
NUM = r'(\d+\s*/\s*\d+|\d+(?:\.\d+)?)'


def parse_qty(raw):
    raw = raw or ''
    s = raw.strip()
    approx = s.startswith('約')
    if approx:
        s = s[1:].strip()
    size = None
    ms = re.match(r'([大中小])\s*(?=\d)', s)   # 「小 1/3コ」のような大きさ指定
    if ms:
        size, s = ms.group(1), s[ms.end():]
    m = re.match(NUM + r'\s*(.*)$', s)
    if not m:
        return {'value': None, 'unit': raw.strip(), 'approx': False, 'raw': raw}
    n = m.group(1).replace(' ', '')
    v = float(Fraction(n)) if '/' in n else float(n)
    out = {'value': v, 'unit': m.group(2).strip(), 'approx': approx, 'raw': raw}
    if size:
        out['size'] = size
    return out


def is_red(cell):
    c = cell.font.color if cell.font else None
    return bool(c is not None and c.type == 'rgb' and str(c.rgb).upper().endswith('FF0000'))


def text(v):
    return '' if v is None else str(v)


def num_fmt(x):
    s = f'{x:.1f}'
    return s[:-2] if s.endswith('.0') else s


def day_total(nuts):
    """その日の食事の栄養価を合計（エネルギーは整数、他は小数1桁）"""
    keys = [('エネルギー', 'kcal'), ('蛋白質', 'g'), ('脂質', 'g'), ('炭水化物', 'g'), ('食塩相当量', 'g')]
    tot = {k: 0.0 for k, _ in keys}
    for n in nuts:
        for k, _ in keys:
            m = re.search(k + r'\s*([\d.]+)', n or '')
            if m:
                tot[k] += float(m.group(1))
    out = []
    for k, u in keys:
        v = f'{round(tot[k])}' if k == 'エネルギー' else f'{tot[k] + 1e-9:.1f}'
        out.append(f'{k} {v}{u}')
    return '\n'.join(out)


def read_sheet(ws):
    m = re.search(r'(\d{4})年(\d{1,2})月(\d{1,2})日', text(ws['B2'].value))
    date = datetime.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    blocks = {}  # (course, meal) -> {'dishes':[], 'nutrition':...}
    cur = None
    dish = None
    for r in range(1, ws.max_row + 1):
        b = text(ws.cell(r, 2).value)
        c = ws.cell(r, 3)
        if b.startswith('（') and b.endswith('）'):
            inner = b[1:-1]
            key = next((v for k, v in COURSE if inner.startswith(k)), None)
            meal = MEAL.get(inner[-1])
            if key and meal:
                cur = blocks.setdefault((key, meal), {'dishes': [], 'nutrition': None})
                dish = None
            continue
        if cur is None:
            continue
        if b in ('材料',):
            continue
        if b.startswith('栄養価'):
            cur['nutrition'] = b.replace('栄養価', '', 1).strip('　 ')
            cur = None
            continue
        if b:
            dish = {'name': b.rstrip(' \u3000'), 'method': text(ws.cell(r, 6).value).rstrip(), 'ingredients': [], 'allergens': []}
            cur['dishes'].append(dish)
            continue
        name = text(c.value)
        if dish is not None and name.strip():
            q1, q3 = text(ws.cell(r, 4).value).strip(), text(ws.cell(r, 5).value).strip()
            ing = {'name': name.rstrip(' \u3000'), 'qty1_raw': q1, 'qty3_raw': q3, 'qty1_parsed': parse_qty(q1)}
            if is_red(c):
                ing['own'] = True
            dish['ingredients'].append(ing)
            alg = text(ws.cell(r, 7).value).strip()
            if alg:   # 料理ごとに、材料のアレルギー表記を空白でつないだ1つの文字列にする
                dish['allergens'] = [' '.join(dish['allergens'] + [alg])]
    return date, blocks


def kondate_photos(repo, monday, key, n_days, n_meals):
    """献立表PDFから盛付写真を取り出し、{(日の番号, 食事の番号): 画像} で返す。
    写真は表の形（横＝日付、縦＝朝・昼・夕）に並んでいるので、ページ上の位置で決める。
    行数・列数が合わなければ None"""
    import fitz, io
    pdf = os.path.join(repo, 'kondate', 'pdf', f'{monday}_{key}.pdf')
    if not os.path.exists(pdf):
        return None
    doc = fitz.open(pdf)
    pg = doc[0]
    cands = []
    for info in pg.get_image_info(xrefs=True):
        x0, y0, x1, y1 = info['bbox']
        w, h = x1 - x0, y1 - y0
        if h > 0 and w > 40 and 1.2 <= w / h <= 1.55:   # 4:3 前後の写真だけ（ロゴ・バナーを除く）
            cands.append(info)
    if not cands:
        return None
    hs = sorted(i['bbox'][3] - i['bbox'][1] for i in cands)
    med = hs[len(hs) // 2]
    cands = [i for i in cands if 0.85 * med <= i['bbox'][3] - i['bbox'][1] <= 1.15 * med]
    cands.sort(key=lambda i: i['bbox'][1])
    rows = []
    for i in cands:   # 上下に半分以上ずれたら別の行
        if rows and abs(i['bbox'][1] - rows[-1][0]['bbox'][1]) < med / 2:
            rows[-1].append(i)
        else:
            rows.append([i])
    if len(rows) != n_meals or any(len(r) != n_days for r in rows):
        print(f'!! {os.path.basename(pdf)}: 写真の並び {[len(r) for r in rows]}（必要 {n_meals}行×{n_days}列）。写真は取り込みません', file=sys.stderr)
        return None
    out = {}
    for mi, r in enumerate(rows):
        for di, info in enumerate(sorted(r, key=lambda i: i['bbox'][0])):
            out[(di, mi)] = Image.open(io.BytesIO(doc.extract_image(info['xref'])['image'])).convert('RGB')
    return out


def save_photo(repo, im):
    w = 420
    im = im.resize((w, round(im.size[1] * w / im.size[0])), Image.LANCZOS)
    tmp = os.path.join(tempfile.mkdtemp(), 'x.jpg')
    im.save(tmp, 'JPEG', quality=82, optimize=True)
    raw = open(tmp, 'rb').read()
    fn = f'photos/{hashlib.sha1(raw).hexdigest()[:12]}.jpg'
    with open(os.path.join(repo, 'recipe', fn), 'wb') as f:
        f.write(raw)
    return fn


def main():
    src, dst = sys.argv[1], sys.argv[2]
    check = '--check' in sys.argv
    s = open(dst, encoding='utf-8').read()
    head = s[:s.index('const COURSES = ')]
    data = json.loads(s[s.index('{'):s.rindex('}') + 1])

    repo = os.path.dirname(os.path.dirname(os.path.abspath(dst)))
    wb = openpyxl.load_workbook(src)
    sheets = [read_sheet(ws) for ws in wb.worksheets if ws.sheet_state == 'visible']
    sheets.sort(key=lambda x: x[0])

    # 献立表PDFからの写真（コース・週ごと）
    photo_of = {}
    if not check:
        weeks = {}
        for date, blocks in sheets:
            weeks.setdefault(date - datetime.timedelta(days=date.weekday()), []).append((date, blocks))
        for monday, days in weeks.items():
            for ck in data:
                meal_names = [m for m in ['朝食', '昼食', '夕食'] if any((ck, m) in b for d, b in days)]
                if not meal_names:
                    continue
                ims = kondate_photos(repo, monday.isoformat(), ck, len(days), len(meal_names))
                if ims:
                    for di, (d, b) in enumerate(days):
                        for mi, m in enumerate(meal_names):
                            photo_of[(ck, d.isoformat(), m)] = save_photo(repo, ims[(di, mi)])

    ng = 0
    for date, blocks in sheets:
        ds = date.isoformat()
        label = f'{date.month}/{date.day}（{WD[date.weekday()]}）'
        for ck in data:
            meals = {}
            for meal in ['朝食', '昼食', '夕食']:
                blk = blocks.get((ck, meal))
                if not blk:
                    continue
                old = data[ck]['days'].get(ds, {}).get('meals', {}).get(meal, {})
                photo = photo_of.get((ck, ds, meal), old.get('photo'))
                meals[meal] = {'dishes': blk['dishes'], 'nutrition': blk['nutrition'], 'photo': photo}
            if not meals:
                print(f'!! {ds} {ck}: Excel にこのコースが見つかりません', file=sys.stderr)
                ng += 1
                continue
            new = {'short_label': label, 'meals': meals,
                   'day_total': day_total([m['nutrition'] for m in meals.values()])}
            if check:
                old = data[ck]['days'].get(ds)
                if old != new:
                    ng += 1
                    for k in new:
                        if old is None or old.get(k) != new[k]:
                            print(f'差分 {ds} {ck} {k}', file=sys.stderr)
                            if k == 'day_total' and old:
                                print('  旧', repr(old.get(k)), '\n  新', repr(new[k]), file=sys.stderr)
            else:
                data[ck]['days'][ds] = new
        print(f'{ds} {label}: ' + ', '.join(f'{c}/{m}' for c, m in sorted(blocks)), file=sys.stderr)

    if check:
        print('一致しました' if ng == 0 else f'{ng} 件の差分', file=sys.stderr)
        sys.exit(1 if ng else 0)
    for ck in data:
        data[ck]['days'] = dict(sorted(data[ck]['days'].items()))
    with open(dst, 'w', encoding='utf-8') as f:
        f.write(head + 'const COURSES = ' + json.dumps(data, ensure_ascii=False, indent=1) + ';\n')
    no_photo = [(ck, ds, m) for ck in data for ds, d in data[ck]['days'].items() for m, md in d['meals'].items() if not md.get('photo')]
    print(f'写真なし: {len(no_photo)} 件', file=sys.stderr)


if __name__ == '__main__':
    main()
